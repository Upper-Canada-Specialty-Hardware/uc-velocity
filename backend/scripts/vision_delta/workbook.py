"""Write a compare's results to delta.xlsx (for people) and delta.json (for tools).

The workbook has a Summary sheet (counts per level, and per level the causes of
its differences), then, when a project review is given, a "projects since"
sheet (one row per recent project, for a person to check) and its "project
changes" detail, then the findings a person resolves (see :data:`FINDINGS`),
then one sheet per level listing every difference.
"""
from __future__ import annotations

import json
import os
from collections import Counter, defaultdict
from datetime import date, datetime

from .compare import REVIEW_COLUMNS, Delta

COLUMNS = ["kind", "project", "context", "field", "vision_value", "velocity_value", "cause", "vision_key", "velocity_key"]
SUMMARY_COLUMNS = ["vision", "velocity", "paired", "same", "differs", "missing", "extra"]


def _cell(v):
    """A value Excel can hold: dates as-is, lists/tuples as text, None blank."""
    if v is None or isinstance(v, (int, float, str, datetime, date)):
        return v
    return str(v)


FINDINGS = [                                               # (sheet, Delta attribute, what the sheet is for)
    ("same job - quotes", "twins", "Vision work orders Velocity lacks, beside Velocity quotes that look alike. "
     "Each is added from Vision with its Vision WO # on it; confirm same job or not."),
    ("same job - POs", "po_twins", "Vision POs Velocity lacks, beside Velocity POs to the same vendor in the same project."),
    ("project numbers", "project_numbers", "Projects both systems numbered on their own after June."),
    ("duplicate projects", "duplicate_projects", "One project entered twice in Velocity since June."),
    ("velocity copies", "copies", "Velocity copies of June quotes; the original's Vision changes are not carried to the copy."),
]


def write_findings(wb, delta: Delta) -> None:
    """Add one sheet per finding a person resolves, each with a notes column.

    Args:
        wb: An open write-only workbook.
        delta: The compare's results.
    """
    for name, attr, about in FINDINGS:
        rows = getattr(delta, attr)
        ws = wb.create_sheet(name[:31])
        ws.freeze_panes = "A3"
        ws.append([about])                                 # first row says what the sheet is for
        cols = list(rows[0]) if rows else []
        ws.append(cols + ["decision / notes"])
        for r in rows:
            ws.append([_cell(r.get(c)) for c in cols] + [None])


def write(delta: Delta, out_dir: str, meta: dict, review: list[dict] | None = None,
          since: date | None = None) -> tuple[str, str]:
    """Write the workbook and the JSON.

    Args:
        delta: The compare's results.
        out_dir: Folder to write into (created if missing).
        meta: Run details for the Summary sheet (inputs, cutoff, time).
        review: Optional :func:`scripts.vision_delta.compare.project_review` rows.
        since: The review's first creation day (names its sheet).

    Returns:
        Paths of ``(delta.xlsx, delta.json)``.
    """
    from openpyxl import Workbook                          # only needed here
    os.makedirs(out_dir, exist_ok=True)
    wb = Workbook(write_only=True)

    # Summary: run details, counts per level, causes per level
    ws = wb.create_sheet("Summary")
    for k, v in meta.items():
        ws.append([k, _cell(v)])
    ws.append([])
    ws.append(["level"] + SUMMARY_COLUMNS)
    for level, c in delta.counts.items():
        ws.append([level] + [c.get(k, 0) for k in SUMMARY_COLUMNS])
    ws.append([])
    ws.append(["level", "kind", "cause", "rows"])
    causes = Counter((r["level"], r["kind"], r["cause"]) for r in delta.rows)
    for (level, kind, why), n in sorted(causes.items()):
        ws.append([level, kind, why, n])

    # Per-project review, then every difference row of those projects
    if review is not None:
        ws = wb.create_sheet(f"projects since {since}"[:31])
        ws.freeze_panes = "B2"                             # keep headings and project number in view
        ws.append(REVIEW_COLUMNS)
        for r in review:
            ws.append([_cell(r[c]) for c in REVIEW_COLUMNS])
        ws = wb.create_sheet("project changes")
        ws.freeze_panes = "A2"
        ws.append(["level"] + COLUMNS)
        order = {r["project"]: i for i, r in enumerate(review)}   # same order as the review sheet
        rows = sorted((r for r in delta.rows if r["project"] in order), key=lambda r: order[r["project"]])
        for r in rows:
            ws.append([r["level"]] + [_cell(r[c]) for c in COLUMNS])

    write_findings(wb, delta)                              # what a person resolves: twins, numbers, duplicates

    # One sheet per level
    by_level = defaultdict(list)
    for r in delta.rows:
        by_level[r["level"]].append(r)
    for level in delta.counts:
        ws = wb.create_sheet(level[:31])                   # Excel's sheet-name limit
        ws.freeze_panes = "A2"
        ws.append(COLUMNS)
        for r in by_level.get(level, []):
            ws.append([_cell(r[c]) for c in COLUMNS])

    xlsx = os.path.join(out_dir, "delta.xlsx")
    wb.save(xlsx)
    js = os.path.join(out_dir, "delta.json")
    with open(js, "w", encoding="utf-8") as f:
        rows = [{k: v for k, v in r.items() if not k.startswith("_")} for r in delta.rows]   # in-memory refs stay out
        json.dump({"meta": meta, "counts": delta.counts, "rows": rows, "project review": review,
                   **{attr: getattr(delta, attr) for _, attr, _ in FINDINGS}}, f, default=str, indent=1)
    return xlsx, js
