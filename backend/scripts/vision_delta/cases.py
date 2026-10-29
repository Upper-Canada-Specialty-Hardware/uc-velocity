"""Build the decisions page's cases from a cutover run, and load them.

A *case* is one thing a person settles: a test project, a duplicate project,
a Vision work order or PO that looks like one already in Velocity, a Velocity
copy of a June quote, a credit-only work order, or one record whose fields
both systems changed. Each case gets a predicted answer and the reason for it.

Two sources:
  * the run's ``vision_sync_changes`` rows the cutover kept for a person
    (actions "held" and "kept: conflict"), grouped into cases here;
  * the findings sheets of the plan's ``report.xlsx`` (same job, duplicate
    projects, Velocity copies), which were reported but never logged.

Loading is insert-only: ``(run_id, kind, case_key)`` is unique and an existing
case is left as it is, so loading twice changes nothing and never touches a
decision already made.

Usage (from ``backend/``)::

    python -m scripts.vision_delta.cases --target <postgres url> --run 36 \\
        --report C:\\dev\\cutover_1029\\plan\\report.xlsx [--commit]

Without ``--commit`` it only prints what it would add.
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from typing import Iterable, Optional

# ---- the steps of the page, in order ------------------------------------- #
# (kind, label, options, what the step is for). The first option of each kind
# is the one that writes nothing except where noted in routes/vision_decisions.
KINDS = [
    ("test project", "Test or empty projects",
     ["left out", "entered by hand"],
     "Vision projects named like test data. The cutover held them, with their work orders."),
    ("duplicate project", "Projects entered twice",
     ["duplicate", "separate projects"],
     "One project entered twice in Velocity since June. Nothing is merged; the answer is recorded for whoever tidies them."),
    ("same job", "Same job in both systems",
     ["separate jobs", "same job"],
     "A Vision work order Velocity lacked, added beside a Velocity quote that looks alike. "
     "'Same job' records the Vision WO # on the Velocity quote."),
    ("same order", "Same PO in both systems",
     ["separate orders", "same order"],
     "A Vision PO Velocity lacked, added beside a Velocity PO to the same vendor. "
     "'Same order' records the Vision PO # on the Velocity PO."),
    ("velocity copy", "Velocity copies of June quotes",
     ["keep both", "copy replaces original"],
     "A Velocity copy of a June quote; Vision's later changes went to the original, not the copy."),
    ("credit memo", "Credit-only work orders",
     ["vision", "velocity", "left out", "entered by hand"],
     "Work orders that hold only credits. June stored them as positive lines; Vision holds the credit."),
    ("conflict", "Changed in both systems",
     ["velocity", "vision", "left out", "entered by hand"],
     "A record both systems changed after June. The cutover kept Velocity's value."),
]
OPTIONS = {k: opts for k, _, opts, _ in KINDS}

CHANGE_ACTIONS = ("held", "kept: conflict")              # log rows a person settles
WO = re.compile(r"\(WO (\d+)\)|^WO (\d+)\b")              # the Vision work order named in a context
FINDING_SHEETS = {                                         # report.xlsx sheet -> kind
    "same job - quotes": "same job",
    "same job - POs": "same order",
    "duplicate projects": "duplicate project",
    "velocity copies": "velocity copy",
}


def _blank(v) -> bool:
    """True for None or text that is only spaces."""
    return v is None or not str(v).strip()


def _wo_of(context: Optional[str], vision_key: Optional[str]) -> Optional[str]:
    """The Vision work order a log row belongs to."""
    m = WO.search(context or "")
    if m:
        return m.group(1) or m.group(2)
    return vision_key


def _row_detail(r: dict) -> dict:
    """One log row as the page shows it: Velocity's value beside Vision's."""
    cause, _, outcome = (r.get("reason") or "").partition(" | ")
    return {"change_id": r["id"], "level": r["level"], "field": r["field"], "record_table": r["record_table"],
            "record_id": r["record_id"], "velocity": r["before"], "vision": r["after"],
            "cause": cause if outcome else None, "context": r["context"]}


# ---- predictions --------------------------------------------------------- #

def predict_conflict(rows: list[dict]) -> tuple[Optional[str], str]:
    """The likely answer for one record both systems changed.

    Args:
        rows: The record's log rows (``before`` = Velocity, ``after`` = Vision).

    Returns:
        ``(option or None, why)``.
    """
    if all(r["field"] is None for r in rows):              # a Vision record Velocity lacks
        return "left out", "Velocity's copy was edited after June without it, so it was probably removed on purpose"
    if any("changed in vision after june" in (r.get("reason") or "") for r in rows):
        return None, "both systems changed it after June; compare the two"
    if all(_blank(r["before"]) and not _blank(r["after"]) for r in rows):
        return "vision", "Velocity is blank and Vision has a value"
    return "velocity", "Velocity was edited after June and Vision records no later change"


def predict_credit(rows: list[dict]) -> tuple[Optional[str], str]:
    """The likely answer for one credit-only work order."""
    if any(r["record_id"] is None for r in rows):          # the work order is not in Velocity at all
        return None, "Velocity never had this credit memo; add it by hand if it is still owed"
    pending = [r for r in rows if r["field"] == "qty_pending" and not _blank(r["after"]) and float(r["after"]) < 0]
    if pending:
        return None, "it holds a credit not yet invoiced, which can be invoiced only once negative quantities ship (#277)"
    return "vision", "Vision holds the processed credit; June stored it as a positive line"


# ---- grouping ------------------------------------------------------------ #

def change_cases(rows: Iterable[dict]) -> list[dict]:
    """Group a run's held and conflict log rows into cases.

    Args:
        rows: ``vision_sync_changes`` rows as dicts (id, level, action,
            record_table, record_id, vision_key, project, context, field,
            before, after, reason).

    Returns:
        Case dicts (kind, case_key, project, title, prediction,
        prediction_reason, detail, change_ids).
    """
    groups: dict[tuple, list[dict]] = defaultdict(list)
    for r in rows:
        reason = r.get("reason") or ""
        if r["action"] == "held" and ("test data" in reason):
            groups[("test project", f"project:{r['project']}")].append(r)       # the project and its held work orders
        elif r["action"] == "held":
            groups[("credit memo", f"wo:{_wo_of(r['context'], r['vision_key'])}")].append(r)
        elif r["record_id"] is not None:
            groups[("conflict", f"{r['record_table']}:{r['record_id']}")].append(r)   # one Velocity record
        else:
            groups[("conflict", f"{r['level']}:vision:{r['vision_key']}")].append(r)  # a Vision record Velocity lacks

    out = []
    for (kind, key), rs in groups.items():
        rs.sort(key=lambda r: r["id"])
        first = rs[0]
        if kind == "test project":
            prediction, why = "left out", "the name reads as test data"
            title = first["context"]
        elif kind == "credit memo":
            prediction, why = predict_credit(rs)
            title = (first["context"] or "").split(" part ")[0].split(" labor ")[0].split(" misc ")[0]
        else:
            prediction, why = predict_conflict(rs)
            title = f"{first['level']}: {first['context']}"
        out.append({"kind": kind, "case_key": key, "project": (first["project"] or "").replace("vision ", "") or None,
                    "title": title, "prediction": prediction, "prediction_reason": why,
                    "detail": {"rows": [_row_detail(r) for r in rs]}, "change_ids": [r["id"] for r in rs]})
    return out


def _sheet_rows(ws) -> list[dict]:
    """A findings sheet's data rows as dicts (row 1 says what it is for, row 2 is the heading)."""
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return []
    head = [h for h in rows[1] if h is not None and h != "decision / notes"]
    return [{h: (str(v) if v is not None else None) for h, v in zip(head, r)} for r in rows[2:] if any(r)]


def _predict_from_text(text: Optional[str], kind: str) -> Optional[str]:
    """Turn a report's prediction text into one of the kind's options."""
    t = (text or "").lower()
    if not t:
        return None
    if kind == "same job":
        return "separate jobs" if "separate" in t else "same job" if "same" in t else None
    if kind == "duplicate project":
        if "duplicate" in t and not re.search(r"\bnot\b", t):   # "not a duplicate" reads the other way
            return "duplicate"
        return "separate projects" if "separate" in t else None
    return None


def finding_cases(workbook_path: str) -> list[dict]:
    """Cases from the report's findings sheets (same job, same PO, duplicate projects, copies).

    Args:
        workbook_path: The plan's ``report.xlsx``.

    Returns:
        Case dicts like :func:`change_cases` (no ``change_ids``).
    """
    from openpyxl import load_workbook
    wb = load_workbook(workbook_path, read_only=True)
    out = []
    for sheet, kind in FINDING_SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        rows = _sheet_rows(wb[sheet])
        if kind == "duplicate project":                    # one case per group of copies
            by_group = defaultdict(list)
            for r in rows:
                by_group[r.get("group")].append(r)
            for group, rs in by_group.items():
                text = next((r.get("prediction") for r in rs if r.get("prediction")), None)
                out.append({"kind": kind, "case_key": f"group:{group}", "project": rs[0].get("uca"),
                            "title": f"{group}: " + ", ".join(r.get("uca") or "?" for r in rs),
                            "prediction": _predict_from_text(text, kind), "prediction_reason": text,
                            "detail": {"found": rs}, "change_ids": []})
            continue
        for r in rows:
            if kind == "same job":
                key, title = f"wo:{r.get('vision WO')}:{r.get('velocity quote id')}", \
                    f"{r.get('project')}: Vision WO {r.get('vision WO')} beside {r.get('velocity quote')}"
                text = r.get("prediction")
            elif kind == "same order":
                key, title = f"po:{r.get('vision PO')}:{r.get('velocity PO')}", \
                    f"{r.get('project')}: Vision PO {r.get('vision PO')} beside {r.get('velocity PO')}"
                text = r.get("evidence")
            else:
                key, title = f"copy:{r.get('original')}:{r.get('copy')}", \
                    f"{r.get('project')}: {r.get('copy')} copies {r.get('original')}"
                text = None
            out.append({"kind": kind, "case_key": key, "project": r.get("project"), "title": title,
                        "prediction": _predict_from_text(text, kind), "prediction_reason": text,
                        "detail": {"found": r}, "change_ids": []})
    return out


# ---- loading ------------------------------------------------------------- #

def load(db, run_id: int, cases: list[dict], commit: bool) -> int:
    """Insert the cases a run does not have yet.

    Args:
        db: A SQLAlchemy session on the target.
        run_id: The cutover run the cases belong to.
        cases: From :func:`change_cases` and :func:`finding_cases`.
        commit: Write when True; otherwise roll back.

    Returns:
        How many cases are new.
    """
    from sqlalchemy.dialects.postgresql import insert
    from models import VisionCase
    have = {(k, key) for k, key in db.query(VisionCase.kind, VisionCase.case_key).filter(VisionCase.run_id == run_id)}
    new = [dict(c, run_id=run_id) for c in cases if (c["kind"], c["case_key"]) not in have]
    for i in range(0, len(new), 500):                      # small statements
        db.execute(insert(VisionCase).values(new[i:i + 500]).on_conflict_do_nothing(
            constraint="uq_vision_cases_run_kind_key"))
    if commit:
        db.commit()
    else:
        db.rollback()
    return len(new)


def main(argv: list[str] | None = None) -> int:
    """Build the run's cases and load them (dry run unless ``--commit``)."""
    p = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    p.add_argument("--target", required=True, help="Postgres URL of the database the run wrote to")
    p.add_argument("--run", type=int, required=True, help="the cutover run id")
    p.add_argument("--report", help="the plan's report.xlsx (same job, duplicate projects, copies)")
    p.add_argument("--commit", action="store_true", help="write the cases; without it nothing is written")
    a = p.parse_args(argv)

    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from models import VisionSyncChange, VisionSyncRun
    with Session(create_engine(a.target)) as db:
        if db.get(VisionSyncRun, a.run) is None:
            print(f"run {a.run} not found; nothing was written")
            return 2
        cols = ("id", "level", "action", "record_table", "record_id", "vision_key", "project", "context",
                "field", "before", "after", "reason")
        q = db.query(*(getattr(VisionSyncChange, c) for c in cols)).filter(
            VisionSyncChange.run_id == a.run, VisionSyncChange.action.in_(CHANGE_ACTIONS))
        cases = change_cases(dict(zip(cols, r)) for r in q)
        if a.report:
            cases += finding_cases(a.report)
        counts = defaultdict(int)
        for c in cases:
            counts[c["kind"]] += 1
        for kind, _, _, _ in KINDS:
            print(f"{kind:20} {counts[kind]:6}")
        n = load(db, a.run, cases, a.commit)
        print(f"{n} new case(s) {'written' if a.commit else 'would be written (dry run; add --commit)'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
