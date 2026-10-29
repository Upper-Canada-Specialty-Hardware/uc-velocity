"""Check a committed cutover run: what it logged is in the database, and nothing is left to apply.

Run from ``backend/`` right after ``python -m scripts.vision_delta.apply --commit``::

    python -m scripts.vision_delta.verify --csv <export folder> --target <postgres url> --run <run id> --out <folder>
                                         [--log-dir <push.py out>/log]

A run made from a plan (plan.py, then push.py) keeps its change log in the log
store, not in ``vision_sync_changes``: pass the local copy push.py kept with
``--log-dir``. Rows on projects the plan left for the archive (no Vision
activity since the plan's ``active_since``) are not expected to be applied.

Four checks, all read-only except the result written onto the run's own row
in ``vision_sync_runs`` (``verify_status``, ``verified_at``, ``verify_detail``):

1. **Same export.** The CSVs in ``--csv`` hash to the fingerprint the run
   recorded, so the compare below sees the data the run applied.
2. **Read-back.** Every change the run logged as added, updated or stamped is
   in the database: an added record exists, a written field holds the value
   logged as "after".
3. **Counts.** For each table, rows after the run minus rows before it equal
   the rows the run logged as added; snapshot tables never shrink; invoices
   are unchanged.
4. **Nothing left.** A fresh compare of the same export with the database,
   put through the update tool's rule, leaves only differences the run logged
   as kept (a conflict, a skip, a failure, information) and only-in-Velocity
   records. A row the rule would still apply, a Vision number still to record,
   or a difference the run never saw, fails the check.

The result is PASS or FAIL, with the detail in ``verify.xlsx`` and ``verify.json``.
"""
from __future__ import annotations

import argparse
import json
import os
import sys
import time
from collections import Counter, defaultdict
from datetime import date, datetime
from typing import Optional
from urllib.parse import urlparse

# Log actions the run wrote to the database, and the ones it only recorded
WRITTEN = ("added", "updated", "stamped")
LEFT = ("skipped", "failed", "kept: conflict", "held", "info")
SNAPSHOT_TABLES = ("quote_snapshots", "po_snapshots")    # grow with every change; not logged one by one

# Compare field -> column, where the log's field is not the column's name
COLUMN = {
    ("projects", "ucsh"): "ucsh_project_number", ("projects", "lead"): "project_lead",
}
SHOWN = 2000                                              # rows per verify.xlsx sheet (the JSON holds all)


def key(level: str, vision_key, velocity_id, field) -> tuple:
    """One difference's identity, the same for a compare row and a log row."""
    vk = None if vision_key in (None, "") else str(vision_key)
    return level, vk, None if velocity_id is None else int(velocity_id), field or None


def same(db_value, logged: Optional[str]) -> bool:
    """Whether a column's value is the one the log says was written.

    The log holds text; numbers compare to the cent (Numeric reads back as
    Decimal), text ignoring case (an enum column read raw holds the member's
    name), a logged day against the stored day.
    """
    if logged is None:
        return db_value is None
    if db_value is None:
        return False
    db_value = getattr(db_value, "value", db_value)       # enum -> its stored text
    if isinstance(db_value, (datetime, date)):
        text = db_value.isoformat(sep=" ") if isinstance(db_value, datetime) else db_value.isoformat()
        return text == logged or (len(logged) == 10 and text[:10] == logged)
    try:
        return abs(float(db_value) - float(logged)) < 0.005
    except (TypeError, ValueError):
        return str(db_value).lower() == logged.lower()     # an enum column holds the member's name ("draft" for "Draft")


def read_back(db, changes: list[dict]) -> list[dict]:
    """Changes the database no longer shows (check 2).

    Args:
        db: An open session.
        changes: The run's logged added / updated / stamped rows.

    Returns:
        One problem row per change that does not read back.
    """
    from sqlalchemy import inspect, text
    columns = {t: {c["name"] for c in inspect(db.get_bind()).get_columns(t)}
               for t in {c["record_table"] for c in changes if c["record_table"]}}
    groups = defaultdict(list)                              # (table, column or None) -> changes
    problems = []
    for c in changes:
        table = c["record_table"]
        if not table or c["record_id"] is None:
            problems.append({**c, "problem": "no record id logged"}); continue
        if c["action"] == "added":
            groups[(table, None)].append(c)                 # an added record: it exists
            continue
        col = COLUMN.get((c["level"], c["field"]), c["field"])
        if c["level"] == "pos" and col == "vendor":
            groups[(table, "vendor")].append(c)             # stored as vendor_id, logged as the name
        elif col in columns[table]:
            groups[(table, col)].append(c)
        else:
            problems.append({**c, "problem": f"column {col} not in {table}"})
    for (table, col), rows in groups.items():
        ids = sorted({r["record_id"] for r in rows})
        # table and column names come from the inspected schema above, never from input
        if col is None:
            sql = f"SELECT id, NULL FROM {table} WHERE id = ANY(:ids)"
        elif col == "vendor":
            sql = ("SELECT po.id, pr.name FROM purchase_orders po LEFT JOIN profiles pr ON pr.id = po.vendor_id "
                   "WHERE po.id = ANY(:ids)")
        else:
            sql = f"SELECT id, {col} FROM {table} WHERE id = ANY(:ids)"
        found = dict(db.execute(text(sql), {"ids": ids}).all())
        for r in rows:
            if r["record_id"] not in found:
                problems.append({**r, "problem": "record not found"})
            elif col is not None and not same(found[r["record_id"]], r["after"]):
                problems.append({**r, "problem": "value differs", "now": str(found[r["record_id"]])})
    return problems


def planned_log(run, log_dir: str) -> list[dict]:
    """A planned run's change log, read from the local copy of the log store.

    Args:
        run: The run's row (``totals`` names every batch and left part with its log key).
        log_dir: The local log folder (push.py's ``<out>/log``).

    Returns:
        Every logged row, shaped like ``vision_sync_changes`` rows.

    Raises:
        SystemExit: A log object is missing from the folder.
    """
    from log_store import FolderStore
    store, out = FolderStore(log_dir), []
    parts = list((run.totals.get("left") or {}).values()) + list((run.totals.get("batches") or {}).values())
    for part in parts:
        try:
            out += [{**e, "run_id": run.id} for e in store.get_json(part["log"])]
        except KeyError:
            raise SystemExit(f"log {part['log']} is not in {log_dir}; run push.py again to fetch it")
    return out


def count_check(run, now: dict, added: Counter) -> list[dict]:
    """Table counts against the logged adds (check 3).

    Args:
        run: The run's row (``counts_before`` / ``counts_after``).
        now: Current counts (staff may have worked since; reported, not judged).
        added: Logged adds per table.

    Returns:
        One row per table, with ``ok``.
    """
    out = []
    before, after = run.counts_before or {}, run.counts_after or {}
    for table in before:
        grew = after.get(table, 0) - before[table]
        if table == "invoices":
            ok = grew == 0                                  # the cutover never touches invoices
        elif table in SNAPSHOT_TABLES:
            ok = grew >= 0
        else:
            ok = grew == added.get(table, 0)
        out.append({"table": table, "before": before[table], "after": after.get(table), "now": now.get(table),
                    "grew": grew, "logged adds": added.get(table, 0) if table not in SNAPSHOT_TABLES else None,
                    "ok": ok})
    return out


def main(argv: list[str] | None = None) -> int:
    """Parse arguments, run the four checks, record the result on the run.

    Args:
        argv: Command-line arguments (defaults to ``sys.argv``).

    Returns:
        0 on PASS, 1 on FAIL, 2 when the run cannot be checked.
    """
    ap = argparse.ArgumentParser(prog="python -m scripts.vision_delta.verify", description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--csv", required=True, help="the Vision CSV export the run applied")
    ap.add_argument("--target", default=os.getenv("VISION_DELTA_TARGET_URL"), help="Velocity Postgres URL")
    ap.add_argument("--run", type=int, required=True, help="the run id apply.py printed")
    ap.add_argument("--out", required=True, help="output folder for verify.xlsx / verify.json")
    ap.add_argument("--log-dir", help="a planned run's change log (push.py's <out>/log, or plan.py's for a local run)")
    args = ap.parse_args(argv)
    if not args.target:
        ap.error("--target (or VISION_DELTA_TARGET_URL) is required")

    os.environ.setdefault("DATABASE_URL", args.target)   # app modules read it at import
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from models import VisionSyncChange, VisionSyncRun
    from . import compare, velocity, vision
    from .apply import APPLIED, INFO, check_schema, fingerprint, rule, stamp_plan, table_counts

    started = time.time()
    print(f"run {args.run} on {urlparse(args.target).hostname}")
    engine = create_engine(args.target)
    detail: dict = {}
    try:
        with Session(engine) as db:
            check_schema(db)
            run = db.get(VisionSyncRun, args.run)
            if run is None or run.mode != "commit":
                print(f"run {args.run} is not a committed run", file=sys.stderr)
                return 2
            planned = (run.totals or {}).get("plan")      # made from a plan: the log is in the store
            if planned:
                if not args.log_dir:
                    print("this run was made from a plan: pass --log-dir", file=sys.stderr)
                    return 2
                log = planned_log(run, args.log_dir)
            else:
                log = [{c.name: getattr(ch, c.name) for c in VisionSyncChange.__table__.columns}
                       for ch in db.query(VisionSyncChange).filter(VisionSyncChange.run_id == run.id)]

            # 1. Same export
            prints = fingerprint(args.csv)
            export_ok = prints == (run.export_fingerprint or {})
            detail["export"] = {"ok": export_ok, "changed files": sorted(
                f for f in set(prints) | set(run.export_fingerprint or {})
                if prints.get(f) != (run.export_fingerprint or {}).get(f))}

            # 2. Read-back
            written = [c for c in log if c["action"] in WRITTEN]
            problems = read_back(db, written)
            detail["read back"] = {"checked": len(written), "problems": len(problems)}

            # 3. Counts
            added = Counter(c["record_table"] for c in log if c["action"] == "added")
            counts = count_check(run, table_counts(db), added)
            detail["counts"] = counts

            # 4. Nothing left: a fresh compare, through the rule, against what the run logged as left
            v = vision.load(args.csv)
            conn = velocity.connect(args.target)
            try:
                t = velocity.load(conn)
            finally:
                conn.close()
            delta = compare.run(v, t, date.fromisoformat(run.june_cutoff))
            archived = lambda label: False                  # a single-transaction run archives nothing
            if planned and (run.totals or {}).get("active_since"):
                from .plan import last_activity
                since, active = date.fromisoformat(run.totals["active_since"]), last_activity(v, delta.project_map)
                archived = lambda label: bool(label) and label in active and active[label] < since
            left = {key(c["level"], c["vision_key"], c["record_id"], c["field"]) for c in log if c["action"] in LEFT}
            still, new, info = [], [], []
            for r in delta.rows:
                if r["kind"] in (compare.EXTRA, compare.COPY):
                    continue                                # only in Velocity: never touched
                if archived(r["project"] or ""):
                    continue                                # left for the archive by the plan
                outcome, reason = rule(r)
                logged = key(r["level"], r["vision_key"], r["velocity_key"], r["field"]) in left
                row = {k: x for k, x in r.items() if not k.startswith("_")} | {"rule": reason}
                if logged:
                    continue                                # the run saw it and left it, as logged
                if outcome == APPLIED:
                    still.append(row)                       # the run should have applied it
                elif outcome == INFO:
                    info.append(row)                        # derived or recorded: reported only
                else:
                    new.append(row)                         # a difference the run never saw
            stamps = [r for r in stamp_plan(delta)[1] if not archived(r["project"] or "")]
            detail["nothing left"] = {"still to apply": len(still), "vision numbers to record": len(stamps),
                                      "new differences": len(new), "new information rows": len(info)}

            ok = export_ok and not problems and all(c["ok"] for c in counts) and not still and not stamps and not new
            run.verify_status, run.verified_at, run.verify_detail = ("pass" if ok else "fail"), datetime.utcnow(), detail
            db.commit()                                     # only the run's own row changes
    finally:
        engine.dispose()

    write_out(args.out, args.run, ok, detail, problems, counts, still, stamps, new, info)
    print(json.dumps(detail, indent=1, default=str))
    print(f"\n{'PASS' if ok else 'FAIL'}  ({time.time() - started:.0f}s)  -> {os.path.join(args.out, 'verify.xlsx')}")
    return 0 if ok else 1


def write_out(out_dir: str, run_id: int, ok: bool, detail: dict, problems: list, counts: list, still: list,
              stamps: list, new: list, info: list) -> None:
    """Write verify.xlsx (a sheet per check) and verify.json (every row)."""
    from openpyxl import Workbook
    from .workbook import _cell
    os.makedirs(out_dir, exist_ok=True)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Summary")
    ws.append(["run", run_id])
    ws.append(["result", "PASS" if ok else "FAIL"])
    ws.append(["checked at", datetime.now().replace(microsecond=0)])
    for check, d in detail.items():
        if isinstance(d, dict):
            for k, val in d.items():
                ws.append([check, k, _cell(val if not isinstance(val, list) else ", ".join(val))])
    for name, rows in (("counts", counts), ("read-back problems", problems), ("still to apply", still),
                       ("vision numbers to record", stamps), ("new differences", new), ("information", info)):
        s = wb.create_sheet(name[:31])
        cols = list(rows[0]) if rows else ["(none)"]
        s.append(cols)
        for r in rows[:SHOWN]:
            s.append([_cell(r.get(c)) for c in cols])
    wb.save(os.path.join(out_dir, "verify.xlsx"))
    with open(os.path.join(out_dir, "verify.json"), "w", encoding="utf-8") as f:
        json.dump({"run": run_id, "result": "pass" if ok else "fail", "detail": detail, "read back problems": problems,
                   "still to apply": still, "vision numbers to record": stamps, "new differences": new,
                   "information": info}, f, default=str, indent=1)


if __name__ == "__main__":
    sys.exit(main())
